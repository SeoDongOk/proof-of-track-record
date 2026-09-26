#!/usr/bin/env python3
"""생성된 .pptx 를 읽어 PNG 로 렌더링한다.

build-deck.py 의 좌표를 다시 쓰지 않고 **파일에 실제로 들어간 것**을 읽는다.
그래야 미리보기와 덱이 어긋나지 않는다.

    python deck/preview.py [pptx] [outdir]   # 기본 deck/preview/slide-NN.png
"""
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont
import os, glob

os.chdir(os.path.dirname(os.path.abspath(__file__)) + '/..')
EMU = 914400
SCALE = 150                      # 1 inch = 150 px
OUT = 'deck/preview'

KR   = '/System/Library/Fonts/Supplemental/AppleGothic.ttf'
KRB  = '/System/Library/Fonts/AppleSDGothicNeo.ttc'
MONO = os.path.expanduser('~/Library/Fonts/D2Coding-Ver1.3.2-20180524.ttc')

_cache = {}
def font(path, px, bold=False):
    key = (path, px, bold)
    if key in _cache: return _cache[key]
    try:
        idx = (6 if bold else 2) if path.endswith('.ttc') and 'AppleSD' in path else 0
        f = ImageFont.truetype(path, px, index=idx)
    except Exception:
        try: f = ImageFont.truetype(path, px)
        except Exception: f = ImageFont.load_default()
    _cache[key] = f
    return f

def wrap(draw, text, f, max_w):
    """픽셀 폭 기준 줄바꿈. 한글은 단어 경계가 드물어 글자 단위로도 자른다."""
    out = []
    for raw in text.split('\n'):
        if not raw: out.append(''); continue
        cur = ''
        for ch in raw:
            if draw.textlength(cur + ch, font=f) <= max_w: cur += ch
            else: out.append(cur); cur = ch
        out.append(cur)
    return out

def render(path='deck/proof-of-track-record.pptx', out=None):
    out = out or OUT
    prs = Presentation(path)
    W = int(prs.slide_width / EMU * SCALE); H = int(prs.slide_height / EMU * SCALE)
    os.makedirs(out, exist_ok=True)
    for old in glob.glob(f'{out}/*.png'): os.remove(old)

    for n, s in enumerate(prs.slides, 1):
        im = Image.new('RGB', (W, H), (12, 17, 27))
        d = ImageDraw.Draw(im)
        for sh in s.shapes:
            x, y = int(sh.left/EMU*SCALE), int(sh.top/EMU*SCALE)
            w, h = int(sh.width/EMU*SCALE), int(sh.height/EMU*SCALE)
            # 이미지
            if sh.shape_type == 13:
                try:
                    pic = Image.open(__import__('io').BytesIO(sh.image.blob)).convert('RGB')
                    im.paste(pic.resize((w, h), Image.LANCZOS), (x, y))
                except Exception as e:
                    d.rectangle([x, y, x+w, y+h], outline=(240,110,110), width=2)
                    d.text((x+8, y+8), f'[이미지 로드 실패 {e}]', fill=(240,110,110))
                continue
            if not sh.has_text_frame: continue
            txt = sh.text_frame.text
            if not txt.strip(): continue
            cy = y
            for p in sh.text_frame.paragraphs:
                if not p.runs: cy += 10; continue
                r0 = p.runs[0]
                pt = r0.font.size.pt if r0.font.size else 18
                px = int(pt / 72 * SCALE)
                name = r0.font.name or ''
                fpath = MONO if name == 'Menlo' else KRB
                f = font(fpath, px, bold=bool(r0.font.bold))
                col = r0.font.color.rgb if r0.font.color and r0.font.color.type is not None else None
                fill = tuple(int(str(col)[i:i+2], 16) for i in (0,2,4)) if col else (236,241,247)
                line = ''.join(r.text for r in p.runs)
                ls = p.line_spacing or 1.25
                for seg in wrap(d, line, f, w):
                    d.text((x, cy), seg, font=f, fill=fill)
                    cy += int(px * ls)
        im.save(f'{out}/slide-{n:02d}.png')
    print(f'  ✅ {out}/slide-01..{n:02d}.png  ({W}x{H})')
    return n

if __name__ == '__main__':
    import sys
    render(*sys.argv[1:3])
