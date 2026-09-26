#!/usr/bin/env python3
"""덱 레이아웃 점검. 슬라이드 밖으로 나가는 것과 서로 겹치는 것을 잡는다.

x 범위까지 보므로 같은 행의 2단 배치를 겹침으로 오인하지 않는다.
이미지와 텍스트가 겹치는 것도 잡는다 — 눈으로만 보다가 놓친 적이 있다.

    python deck/check.py [pptx]
"""
from pptx import Presentation
import os, sys
os.chdir(os.path.dirname(os.path.abspath(__file__)) + '/..')
E = 914400

def est_height(sh):
    """텍스트 박스의 실제 높이 추정. 폭 대비 줄바꿈을 계산한다."""
    t = sh.text_frame.text
    w = sh.width / E
    sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
    pt = max(sizes) if sizes else 18
    mono = any(r.font.name == 'Menlo' for p in sh.text_frame.paragraphs for r in p.runs)
    cpl = max(1, int(w * 72 / (pt * (0.62 if mono else 0.80))))
    lines = sum(max(1, -(-len(l) // cpl)) for l in t.split('\n'))
    return lines * pt * 1.30 / 72, pt

def boxes(s):
    out = []
    for sh in s.shapes:
        x, y = sh.left / E, sh.top / E
        w, h = sh.width / E, sh.height / E
        if sh.shape_type == 13:
            out.append(('img', x, y, w, h, '이미지'))
        elif sh.has_text_frame and sh.text_frame.text.strip():
            eh, pt = est_height(sh)
            # 배경 도형은 제외
            if w > 13 and h > 7: continue
            out.append(('txt', x, y, w, eh, sh.text_frame.text.split('\n')[0][:30]))
    return out

def overlap(a, b, pad=0.05):
    _, ax, ay, aw, ah, _ = a; _, bx, by, bw, bh, _ = b
    return (ax < bx + bw - pad and bx < ax + aw - pad and
            ay < by + bh - pad and by < ay + ah - pad)

prs = Presentation(sys.argv[1] if len(sys.argv) > 1 else 'deck/proof-of-track-record.pptx')
SW, SH = prs.slide_width / E, prs.slide_height / E
issues = 0
for i, s in enumerate(prs.slides, 1):
    bs = boxes(s)
    for b in bs:
        kind, x, y, w, h, label = b
        if y + h > SH - 0.06 or x + w > SW - 0.06:
            print(f'  [{i:2}] ❌ 슬라이드 밖  {kind} {label!r}  하단 {y+h:.2f}/{SH:.2f}  우측 {x+w:.2f}/{SW:.2f}')
            issues += 1
    for j in range(len(bs)):
        for k in range(j + 1, len(bs)):
            if overlap(bs[j], bs[k]):
                a, b2 = bs[j], bs[k]
                print(f'  [{i:2}] ❌ 겹침  {a[0]} {a[5]!r} (y {a[2]:.2f}~{a[2]+a[4]:.2f})'
                      f'  ↔  {b2[0]} {b2[5]!r} (y {b2[2]:.2f}~{b2[2]+b2[4]:.2f})')
                issues += 1
print(f'\n  {len(prs.slides._sldIdLst)}장 / 문제 {issues}건')
sys.exit(1 if issues else 0)
