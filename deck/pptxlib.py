"""덱 조립 공용 헬퍼. build-deck.py(읽기용) 와 build-talk.py(발표용) 가 같이 쓴다.
색·폰트·좌표 헬퍼를 한 곳에 두어 두 덱의 생김새가 어긋나지 않게 한다.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

os.chdir(os.path.dirname(os.path.abspath(__file__)) + '/..')

BG  = RGBColor(0x0C,0x11,0x1B); FG  = RGBColor(0xEC,0xF1,0xF7)
DIM = RGBColor(0x8B,0x9A,0xAF); ACC = RGBColor(0x5E,0xE0,0xB0)
RED = RGBColor(0xF0,0x6E,0x6E); AMB = RGBColor(0xE8,0xB4,0x5C)
KR  = 'Apple SD Gothic Neo'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

def slide():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG; bg.line.fill.background()
    bg.shadow.inherit = False
    return s

def text(s, txt, x, y, w, h, size=20, color=FG, bold=False, align=PP_ALIGN.LEFT, sp=1.25, font=KR):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(txt.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.alignment = align; p.line_spacing = sp
        for r in p.runs:
            r.font.size = Pt(size); r.font.color.rgb = color
            r.font.bold = bold; r.font.name = font
    return tb

mono = lambda s, t, x, y, w, h, size=14, color=ACC: text(s, t, x, y, w, h, size, color, sp=1.35, font='Menlo')
tag  = lambda s, t, x, y, color=DIM, size=13: text(s, t, x, y, 8, 0.35, size=size, color=color)
pic  = lambda s, p, x, y, w: s.shapes.add_picture(p, Inches(x), Inches(y), width=Inches(w))

def pic_fit(s, path, y, max_h, sw=13.333):
    """높이에 맞춰 넣고 가로 중앙 정렬. 폭 기준으로 넣으면 세로가 넘치는 이미지용."""
    from PIL import Image
    iw, ih = Image.open(path).size
    w = max_h * iw / ih
    if w > sw - 0.9: w = sw - 0.9; max_h = w * ih / iw
    return s.shapes.add_picture(path, Inches((sw - w) / 2), Inches(y), width=Inches(w))

def save(path):
    prs.save(path)
    print(f'  ✅ {path}  ({len(prs.slides._sldIdLst)}장)')
